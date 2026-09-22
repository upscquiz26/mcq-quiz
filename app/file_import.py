"""
Importing questions from other kinds of file: spreadsheets (CSV / XLSX), Word documents (DOCX) and pictures (PNG / JPG …).

  * CSV / XLSX — one question per row. The admin says which column holds which field (the app guesses from the headers, and remembers the
    choice for the next file with the same headers). Answers must be the letters A–D; 1–4 is refused, never guessed.
  * DOCX — read with the same question cutter as text PDFs (numbered questions, options (a)–(d), optional "Answer:" lines and a key block at
    the end). Word's *automatic* numbering isn't stored in the text, so typed numbers ("1.") are needed.
  * Images — one or more page pictures are put together into a PDF and read by OCR exactly like a scanned paper (see ingest.py).

Everything lands as "needs review". Nothing is confirmed, nothing is published, and there is no original page to compare with, so the review
screen shows only the text — check the source yourself.
"""
import csv
import hashlib
import io
import json
import re
from dataclasses import dataclass, field
from datetime import datetime

from app import models
from app.models import QStatus

TABLE_EXTS = (".csv", ".xlsx", ".xlsm")
DOCX_EXTS = (".docx",)
IMAGE_EXTS = (".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff")
LETTERS = ("A", "B", "C", "D")
SHOWN_PROBLEMS = 25            # how many individual row problems are listed before "… and N more"

# (field, label, required)
FIELDS = [
    ("number", "Question number", False), ("question", "Question", True),
    ("option_a", "Option A", True), ("option_b", "Option B", True), ("option_c", "Option C", True), ("option_d", "Option D", True),
    ("answer", "Correct answer (A–D)", False), ("explanation", "Explanation", False), ("subject", "Subject", False), ("topic", "Topic", False),
]
FIELD_NAMES = [f for f, _, _ in FIELDS]


class FileImportError(Exception):
    """The file can't be read; str(e) says why, in words for the admin."""


def kind_of(filename: str) -> str | None:
    name = (filename or "").lower()
    for kind, exts in (("table", TABLE_EXTS), ("docx", DOCX_EXTS), ("images", IMAGE_EXTS)):
        if name.endswith(exts):
            return kind
    return None


# --------------------------------------------------------------------------- spreadsheets

@dataclass
class Table:
    headers: list[str]
    rows: list[list[str]]
    sheets: list[str] = field(default_factory=list)
    sheet: str | None = None
    note: str | None = None


def _clean(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)                                   # Excel stores 3 as 3.0
    return re.sub(r"[ \t\r\f\v ]+", " ", str(value)).strip()


def _table_from(grid: list[list[str]], sheets, sheet, note) -> Table:
    grid = [row for row in grid if any(cell for cell in row)]
    if len(grid) < 2:
        raise FileImportError("The file needs a header row and at least one row of questions.")
    width = max(len(r) for r in grid)
    grid = [r + [""] * (width - len(r)) for r in grid]
    headers = [h or f"Column {i + 1}" for i, h in enumerate(grid[0])]
    return Table(headers=headers, rows=grid[1:], sheets=sheets or [], sheet=sheet, note=note)


def read_table(data: bytes, filename: str, sheet: str | None = None) -> Table:
    name = (filename or "").lower()
    if name.endswith(".csv"):
        text, note = None, None
        for encoding in ("utf-8-sig", "cp1252"):
            try:
                text = data.decode(encoding)
                note = None if encoding == "utf-8-sig" else "Read as Windows-1252 (it isn't UTF-8); check accented characters."
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise FileImportError("This CSV can't be decoded. Save it as UTF-8 from your spreadsheet program and try again.")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        try:
            grid = [[_clean(c) for c in row] for row in csv.reader(io.StringIO(text), dialect)]
        except csv.Error as e:
            raise FileImportError(f"This CSV can't be read ({e}).")
        return _table_from(grid, [], None, note)
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:
        raise FileImportError(f"This spreadsheet can't be opened ({type(e).__name__}). Is it a real .xlsx file?")
    try:
        sheets = list(wb.sheetnames)
        chosen = sheet if sheet in sheets else sheets[0]
        grid = [[_clean(c) for c in row] for row in wb[chosen].iter_rows(values_only=True)]
    finally:
        wb.close()
    note = f"The workbook has {len(sheets)} sheets; this is “{chosen}”." if len(sheets) > 1 else None
    return _table_from(grid, sheets, chosen, note)


def _norm(header: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", (header or "").lower())


def signature(headers: list[str]) -> str:
    return hashlib.sha1("|".join(sorted(n for n in (_norm(h) for h in headers) if n)).encode()).hexdigest()[:16]


SYNONYMS = {
    "number": ("no", "number", "qno", "qn", "qnum", "questionno", "questionnumber", "sno", "srno", "slno", "sl", "serial", "serialno", "id"),
    "question": ("question", "questions", "q", "stem", "questiontext", "text", "prompt", "statement"),
    "option_a": ("optiona", "opta", "a", "choicea", "option1", "opt1", "choice1", "answera"),
    "option_b": ("optionb", "optb", "b", "choiceb", "option2", "opt2", "choice2", "answerb"),
    "option_c": ("optionc", "optc", "c", "choicec", "option3", "opt3", "choice3", "answerc"),
    "option_d": ("optiond", "optd", "d", "choiced", "option4", "opt4", "choice4", "answerd"),
    "answer": ("answer", "ans", "correctanswer", "correctoption", "correct", "key", "answerkey", "rightanswer", "correctchoice"),
    "explanation": ("explanation", "explanations", "solution", "rationale", "reason", "description"),
    "subject": ("subject", "subj", "section"),
    "topic": ("topic", "subtopic", "chapter"),
}
CONTAINS = {"answer": ("answer", "correct"), "explanation": ("explanation", "solution"), "subject": ("subject",), "topic": ("topic",)}


def guess_mapping(headers: list[str]) -> dict[str, int]:
    """{field: column index} guessed from header names. A column is used for one field only."""
    normal = [_norm(h) for h in headers]
    mapping: dict[str, int] = {}
    used: set[int] = set()
    for name in FIELD_NAMES:
        for i, n in enumerate(normal):
            if i not in used and n in SYNONYMS[name]:
                mapping[name] = i
                used.add(i)
                break
    for name, words in CONTAINS.items():
        if name in mapping:
            continue
        for i, n in enumerate(normal):
            if i not in used and any(w in n for w in words):
                mapping[name] = i
                used.add(i)
                break
    if "question" not in mapping:
        for i, n in enumerate(normal):
            if i not in used and n.startswith("question") and not n.endswith(("no", "number", "num")):
                mapping["question"] = i
                used.add(i)
                break
    return mapping


def recall_mapping(db, headers: list[str]) -> tuple[dict[str, int], models.ImportMapping | None]:
    """The mapping remembered for a file with these headers (as column indexes), or ({}, None)."""
    row = db.query(models.ImportMapping).filter_by(signature=signature(headers)).first()
    if not row:
        return {}, None
    by_name = json.loads(row.mapping_json)
    index = {h: i for i, h in enumerate(headers)}
    return {f: index[h] for f, h in by_name.items() if h in index and f in FIELD_NAMES}, row


def remember_mapping(db, headers: list[str], mapping: dict[str, int]) -> None:
    by_name = {f: headers[i] for f, i in mapping.items() if 0 <= i < len(headers)}
    sig = signature(headers)
    row = db.query(models.ImportMapping).filter_by(signature=sig).first()
    if row:
        row.mapping_json, row.headers_json, row.uses, row.updated_at = json.dumps(by_name), json.dumps(headers), row.uses + 1, datetime.utcnow()
    else:
        db.add(models.ImportMapping(signature=sig, headers_json=json.dumps(headers), mapping_json=json.dumps(by_name)))


def read_mapping_form(form, width: int) -> dict[str, int]:
    """{field: column index} from posted fields named map_<field> holding a column number (or empty for 'none')."""
    mapping = {}
    for name in FIELD_NAMES:
        value = str(form.get("map_" + name, "")).strip()
        if value.isdigit() and int(value) < width:
            mapping[name] = int(value)
    return mapping


def normalise_answer(value: str) -> str | None:
    """'b', 'B', '(b)', 'B.', 'Option B', 'answer: b' -> 'B'. Anything else (1-4, two letters, words) -> None: never guessed."""
    text = (value or "").strip()
    m = re.fullmatch(r"(?:(?:option|answer|ans|choice)\s*[:\-–.]?\s*)?[(\[]?\s*([a-dA-D])\s*[)\].:]?", text, re.I)
    return m.group(1).upper() if m else None


# --------------------------------------------------------------------------- what a parsed file looks like

@dataclass
class Parsed:
    questions: list[dict] = field(default_factory=list)      # number, text, options[4], answer, explanation, subject, topic, flags, hash
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    answer_source: str = "file"
    method: str = "csv"

    @property
    def ok(self) -> bool:
        return bool(self.questions) and not self.errors


def _listed(items: list, limit: int = SHOWN_PROBLEMS) -> str:
    shown = ", ".join(str(i) for i in items[:limit])
    return shown + (f" … and {len(items) - limit} more" if len(items) > limit else "")


def _finish(parsed: Parsed) -> Parsed:
    from app import json_import
    from app.ocr_extract import sanity_flags
    from app.text_extract import DEVANAGARI
    for q in parsed.questions:
        flags = list(q.get("flags", []))
        for extra in sanity_flags(q["text"], q["options"]) + (["hindi_text"] if DEVANAGARI.search(q["text"] + " " + " ".join(q["options"])) else []):
            if extra not in flags:
                flags.append(extra)
        q["flags"] = flags
        q["hash"] = json_import.norm_hash(q["text"], q["options"])
    return parsed


def parse_rows(table: Table, mapping: dict[str, int], subject_names: list[str], topics: dict[str, set[str]] | None = None,
               method: str = "csv") -> Parsed:
    """Turn spreadsheet rows into questions, or say exactly which rows are wrong. Errors block the import; warnings don't."""
    parsed = Parsed(method=method, answer_source="file")
    missing = [label for f, label, required in FIELDS if required and f not in mapping]
    if missing:
        parsed.errors.append("Choose a column for: " + ", ".join(missing) + ".")
        return parsed
    chosen = [i for i in mapping.values()]
    if len(chosen) != len(set(chosen)):
        parsed.errors.append("The same column is chosen for two different fields.")
        return parsed
    by_lower = {s.lower(): s for s in subject_names}
    bad_numbers, no_text, bad_options, duplicates_seen = [], [], [], []
    seen: dict[int, int] = {}
    unreadable_answers, no_answers, unknown_subjects, unknown_topics = [], [], set(), 0
    counter = 0
    for offset, row in enumerate(table.rows):
        line = offset + 2                                       # the header is line 1 (blank lines are skipped, so this is approximate)
        counter += 1
        cell = lambda f: row[mapping[f]].strip() if f in mapping else ""      # noqa: E731
        number = counter
        if "number" in mapping:
            m = re.search(r"\d+", cell("number"))
            if not m:
                bad_numbers.append(line)
                continue
            number = int(m.group())
        if number in seen:
            duplicates_seen.append(f"{number} (lines {seen[number]} and {line})")
            continue
        seen[number] = line
        text = cell("question")
        options = [cell(f) for f in ("option_a", "option_b", "option_c", "option_d")]
        if not text:
            no_text.append(line)
            continue
        empty = [LETTERS[i] for i, o in enumerate(options) if not o]
        if empty:
            bad_options.append(f"line {line}: option {', '.join(empty)} is empty")
            continue
        flags = []
        answer = None
        raw_answer = cell("answer")
        if raw_answer:
            answer = normalise_answer(raw_answer)
            if answer is None:
                unreadable_answers.append(f"Q{number} “{raw_answer[:12]}”")
                flags.append("answer_unclear")
        elif "answer" in mapping:
            no_answers.append(number)
            flags.append("no_answer_found")
        else:
            flags.append("no_answer_found")
        subject = None
        if cell("subject"):
            subject = by_lower.get(cell("subject").lower())
            if subject is None:
                unknown_subjects.add(cell("subject")[:30])
        topic = None
        if cell("topic"):
            if subject and cell("topic").lower() in (topics or {}).get(subject, set()):
                topic = cell("topic")
            else:
                unknown_topics += 1
        parsed.questions.append({"number": number, "text": text, "options": options, "answer": answer,
                                 "explanation": cell("explanation") or None, "subject": subject, "topic": topic, "flags": flags})
    if bad_numbers:
        parsed.errors.append(f"No readable question number on line{'s' if len(bad_numbers) != 1 else ''} {_listed(bad_numbers)}.")
    if duplicates_seen:
        parsed.errors.append("Question number used twice: " + _listed(duplicates_seen) + ".")
    if no_text:
        parsed.errors.append(f"No question text on line{'s' if len(no_text) != 1 else ''} {_listed(no_text)}.")
    if bad_options:
        parsed.errors.append("Missing options — " + "; ".join(bad_options[:SHOWN_PROBLEMS]) + (" …" if len(bad_options) > SHOWN_PROBLEMS else "") + ".")
    if not parsed.questions and not parsed.errors:
        parsed.errors.append("No questions were found in the file.")
    if unreadable_answers:
        parsed.warnings.append(f"{len(unreadable_answers)} answer{'s' if len(unreadable_answers) != 1 else ''} can't be read as A, B, C or D "
                               f"({_listed(unreadable_answers, 8)}). They were left blank and flagged; digits like 1–4 are never guessed as letters.")
    if "answer" not in mapping:
        parsed.warnings.append("No answer column was chosen, so every question needs its answer set on the review page (or add a key afterwards).")
    elif no_answers:
        parsed.warnings.append(f"No answer for question{'s' if len(no_answers) != 1 else ''} {_listed(no_answers)}.")
    if unknown_subjects:
        parsed.warnings.append("Subjects not in the fixed list were left blank: " + ", ".join(sorted(unknown_subjects)[:8]) + ".")
    if unknown_topics:
        parsed.warnings.append(f"{unknown_topics} topic{'s' if unknown_topics != 1 else ''} weren't used (topics are only taken when they already exist under the subject).")
    numbers = sorted(seen)
    gaps = sorted(set(range(1, numbers[-1] + 1)) - set(numbers)) if numbers and "number" in mapping else []
    if gaps:
        parsed.warnings.append(f"Question numbers missing from the file: {_listed(gaps)}.")
    return _finish(parsed)


# --------------------------------------------------------------------------- Word documents

def parse_docx(data: bytes) -> Parsed:
    try:
        import docx
        from docx.table import Table as DocxTable
        from docx.text.paragraph import Paragraph
        document = docx.Document(io.BytesIO(data))
    except Exception as e:
        raise FileImportError(f"This document can't be opened ({type(e).__name__}). Is it a real .docx file (not an old .doc)?")
    from app import key_parse, text_extract

    lines: list[dict] = []
    automatic = 0
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = Paragraph(child, document)
            if paragraph._p.pPr is not None and paragraph._p.pPr.numPr is not None:
                automatic += 1
            chunks = paragraph.text.replace(" ", " ").splitlines()
        elif tag == "tbl":
            chunks = ["  ".join(_clean(c.text) for c in row.cells if _clean(c.text)) for row in DocxTable(child, document).rows]
        else:
            continue
        for chunk in chunks:
            chunk = re.sub(r"[ \t]+", " ", chunk).strip()
            if chunk:
                lines.append({"text": chunk, "page": 0})

    parsed = Parsed(method="docx", answer_source="inline")
    lines, key_block = text_extract.split_key_block(lines)
    raw, inline = text_extract.cut_questions(lines)
    for item in raw:
        q = text_extract._question_from(item, None)
        flags = list(q["flags"])
        if inline and not q["answer_seen"]:
            flags.append("no_answer_found")
        elif inline and q["answer"] is None:
            flags.append("answer_unclear")
        parsed.questions.append({"number": q["question_number"], "text": q["text"],
                                 "options": [q["option_a"], q["option_b"], q["option_c"], q["option_d"]], "answer": q["answer"],
                                 "explanation": q["explanation"], "subject": None, "topic": None, "flags": flags})
    if not parsed.questions:
        parsed.errors.append("No questions were found in the document. Questions must be numbered “1.”, “2.” … as typed text, with options (a)–(d).")
    else:
        numbers = {q["number"] for q in parsed.questions}
        gaps = sorted(set(range(1, max(numbers) + 1)) - numbers)
        if gaps:
            parsed.warnings.append(f"Question numbers not detected: {_listed(gaps)}.")
        merged = [q["number"] for q in parsed.questions if "maybe_merged" in q["flags"]]
        if merged:
            parsed.warnings.append("Question(s) " + _listed(merged, 15) + " look like two questions merged into one — split them on the review page.")
        incomplete = [q["number"] for q in parsed.questions if "options_not_found" in q["flags"]]
        if incomplete:
            parsed.warnings.append("Options (a)–(d) weren't found for question(s) " + _listed(incomplete, 15) + ". Later questions may have been merged into them.")
        if not inline and key_block:
            key = key_parse.parse_key_text(key_block, None)
            if key.answers and not key.errors:
                filled = 0
                for q in parsed.questions:
                    letter = key.answers.get(q["number"])
                    if letter and not q["answer"]:
                        q["answer"], filled = letter, filled + 1
                        q["flags"] = [f for f in q["flags"] if f != "no_answer_found"]
                if filled:
                    parsed.answer_source = "paper_end"
                    parsed.warnings.append(f"An answer key at the end of the document gave {filled} answers.")
            else:
                reason = "; ".join(i.message for i in key.errors[:2])
                parsed.warnings.append("An answer key block was found at the end of the document but couldn't be used"
                                       + (f" ({reason})." if reason else "."))
        for q in parsed.questions:
            if not q["answer"] and not {"no_answer_found", "answer_unclear"} & set(q["flags"]):
                q["flags"].append("no_answer_found")
        if not any(q["answer"] for q in parsed.questions):
            parsed.warnings.append("The document has no answers, so every question needs its answer set on the review page.")
    if automatic >= 3 and (not parsed.questions or any("Question numbers not detected" in w for w in parsed.warnings)):
        parsed.warnings.append(f"{automatic} paragraphs use Word's automatic numbering. Those numbers aren't stored in the text, so numbered questions may be missing: "
                               "type the numbers (1. 2. 3. …) into the document and import it again.")
    pictures = len(document.inline_shapes)
    if pictures:
        parsed.warnings.append(f"The document contains {pictures} picture{'s' if pictures != 1 else ''}; pictures aren't imported. "
                               "A question that depends on one needs its picture checked in the original.")
    return _finish(parsed)


# --------------------------------------------------------------------------- pictures

def images_to_pdf(images: list[bytes], out_path: str) -> int:
    """Puts page pictures, in the order given, into one PDF (each picture is a page). Returns the number of pages."""
    from PIL import Image, UnidentifiedImageError
    pages = []
    for i, data in enumerate(images, start=1):
        try:
            image = Image.open(io.BytesIO(data))
            image.load()
        except (UnidentifiedImageError, OSError):
            raise FileImportError(f"Picture {i} can't be read as an image.")
        pages.append(image.convert("RGB"))
    if not pages:
        raise FileImportError("Choose at least one picture.")
    pages[0].save(out_path, "PDF", save_all=len(pages) > 1, append_images=pages[1:], resolution=200.0)
    return len(pages)


# --------------------------------------------------------------------------- saving

def save_questions(db, user, parsed: Parsed, paper: models.Paper) -> dict:
    """Writes the parsed questions into `paper` as needs-review questions (the caller commits). Numbers the paper already has are skipped.
    Afterwards subject hints and duplicate checks are run for the paper. Returns counts."""
    from app import duplicates, subject_hints
    subjects = {s.name: s.id for s in db.query(models.Subject).all()}
    topic_ids = {(t.subject_id, t.name.lower()): t.id for t in db.query(models.Topic).all()}
    present = {q.question_number for q in db.query(models.Question).filter_by(paper_id=paper.id)}
    counts = {"created": 0, "skipped_existing": 0}
    for item in parsed.questions:
        if item["number"] in present:
            counts["skipped_existing"] += 1
            continue
        subject_id = subjects.get(item["subject"]) if item["subject"] else None
        topic_id = topic_ids.get((subject_id, item["topic"].lower())) if item["topic"] and subject_id else None
        answer_source = None
        if item["answer"]:
            answer_source = parsed.answer_source
        db.add(models.Question(
            paper_id=paper.id, question_number=item["number"], text=item["text"], option_a=item["options"][0], option_b=item["options"][1],
            option_c=item["options"][2], option_d=item["options"][3], correct_answer=item["answer"], answer_source=answer_source,
            explanation=item["explanation"], explanation_status="unverified" if item["explanation"] else None,
            subject_id=subject_id, topic_id=topic_id, ocr_flags=",".join(item["flags"]) or None, norm_hash=item["hash"],
            status=QStatus.NEEDS_REVIEW, source=parsed.method, extraction_method=parsed.method,
        ))
        counts["created"] += 1
    db.flush()
    subject_hints.suggest_for_paper(db, paper.id)
    counts["duplicates"] = duplicates.scan(db, paper.id)
    return counts
