"""Import stage 9: CSV / XLSX with remembered column mapping, DOCX, and picture imports."""
import csv
import io
import itertools
import json
import re
import uuid

import pytest
from PIL import Image

from app import file_import as fi
from app import ingest, models, ocr_extract
from app.models import QStatus
from app.practice import attempts
from conftest import make_student_client

_run = itertools.count(1)
HEADERS = ["Q No", "Question", "Option A", "Option B", "Option C", "Option D", "Answer", "Explanation", "Subject"]
SUBJECTS = ["Polity", "History", "Geography", "Economy", "Environment", "Science & Tech", "Current Affairs", "CSAT", "Other"]


def stem():
    return f"Which of the following statements about the charter {uuid.uuid4().hex[:10]} of the {uuid.uuid4().hex[:8]} council is correct?"


def rows_of(n=3, answer="B", subject="Polity"):
    return [[str(i), stem(), f"alpha {i}", f"beta {i}", f"gamma {i}", f"delta {i}", answer, f"Because {i}.", subject] for i in range(1, n + 1)]


def csv_bytes(rows, headers=HEADERS, delimiter=",", encoding="utf-8-sig"):
    out = io.StringIO()
    writer = csv.writer(out, delimiter=delimiter)
    writer.writerow(headers)
    writer.writerows(rows)
    return out.getvalue().encode(encoding)


def xlsx_bytes(sheets: dict):
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def docx_bytes(paragraphs, tables=(), picture=False, numbered=0):
    import docx
    d = docx.Document()
    for text in paragraphs:
        d.add_paragraph(text)
    for _ in range(numbered):
        p = d.add_paragraph("An automatically numbered paragraph")
        p._p.get_or_add_pPr().get_or_add_numPr()
    for rows in tables:
        t = d.add_table(rows=len(rows), cols=len(rows[0]))
        for i, row in enumerate(rows):
            for j, cell in enumerate(row):
                t.cell(i, j).text = cell
    if picture:
        img = io.BytesIO()
        Image.new("RGB", (20, 20), "white").save(img, "PNG")
        img.seek(0)
        d.add_picture(img)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def docx_questions(n=3, answers=True):
    lines = []
    for i in range(1, n + 1):
        lines += [f"{i}. {stem()}", "(a) alpha (b) beta", "(c) gamma (d) delta"]
        if answers:
            lines.append(f"Answer: ({'abcd'[(i - 1) % 4]})")
    return lines


def png(color="white", size=(300, 200)):
    buf = io.BytesIO()
    Image.new("RGB", size, color).save(buf, "PNG")
    return buf.getvalue()


# --------------------------------------------------------------------------- reading spreadsheets

def test_files_are_told_apart_by_extension():
    assert fi.kind_of("a.CSV") == "table" and fi.kind_of("b.xlsx") == "table" and fi.kind_of("c.docx") == "docx"
    assert fi.kind_of("d.JPG") == "images" and fi.kind_of("e.pdf") is None and fi.kind_of("") is None


def test_csv_reads_with_a_bom_semicolons_and_blank_lines():
    data = csv_bytes([["1", "Q one?", "a", "b", "c", "d"], ["", "", "", "", "", ""], ["2", "Q two?", "a", "b", "c", "d"]],
                     headers=["No", "Question", "A", "B", "C", "D"], delimiter=";")
    table = fi.read_table(data, "q.csv")
    assert table.headers == ["No", "Question", "A", "B", "C", "D"] and len(table.rows) == 2 and table.rows[1][1] == "Q two?"


def test_a_windows_1252_csv_is_read_and_says_so():
    table = fi.read_table("No,Question\n1,Caf\xe9?\n".encode("cp1252"), "q.csv")
    assert table.rows[0][1] == "Café?" and "Windows-1252" in table.note


def test_a_file_without_data_rows_is_refused():
    with pytest.raises(fi.FileImportError):
        fi.read_table(b"No,Question\n", "q.csv")
    with pytest.raises(fi.FileImportError):
        fi.read_table(b"not a spreadsheet at all", "q.xlsx")


def test_ragged_rows_are_padded_and_empty_headers_named():
    table = fi.read_table(b"No,,Question\n1,x\n2,y,Q?\n", "q.csv")
    assert table.headers == ["No", "Column 2", "Question"] and table.rows[0] == ["1", "x", ""]


def test_xlsx_reads_numbers_without_decimals_and_lets_you_pick_a_sheet():
    data = xlsx_bytes({"First": [["No", "Question"], [1, "Q one?"], [2.0, "Q two?"]], "Second": [["No", "Question"], [7, "Other?"]]})
    table = fi.read_table(data, "q.xlsx")
    assert table.sheets == ["First", "Second"] and table.sheet == "First" and table.rows == [["1", "Q one?"], ["2", "Q two?"]]
    assert "2 sheets" in table.note
    assert fi.read_table(data, "q.xlsx", "Second").rows == [["7", "Other?"]]
    assert fi.read_table(data, "q.xlsx", "Nope").sheet == "First"                                   # an unknown sheet falls back to the first


# --------------------------------------------------------------------------- the column mapping

def test_the_mapping_is_guessed_from_common_headers():
    m = fi.guess_mapping(["S.No", "Question", "Option A", "Option B", "Option C", "Option D", "Correct Answer", "Solution", "Subject", "Topic"])
    assert m == {"number": 0, "question": 1, "option_a": 2, "option_b": 3, "option_c": 4, "option_d": 5, "answer": 6, "explanation": 7,
                 "subject": 8, "topic": 9}
    m = fi.guess_mapping(["Q", "Stem", "1", "2"])
    assert m["question"] in (0, 1) and "option_a" not in m
    m = fi.guess_mapping(["Question No", "Question Text", "Opt1", "Opt2", "Opt3", "Opt4", "Ans"])
    assert m["number"] == 0 and m["question"] == 1 and m["option_d"] == 5 and m["answer"] == 6
    assert len(set(m.values())) == len(m)                                                          # a column is used for one field only


def test_the_signature_ignores_order_case_and_punctuation():
    assert fi.signature(["Q No", "Question", "Option A"]) == fi.signature(["option a", "QUESTION", "q.no"]) != fi.signature(["Q No", "Question"])


def test_a_remembered_mapping_follows_the_header_names_not_the_column_positions(db):
    headers = [f"H{uuid.uuid4().hex[:6]}" for _ in range(4)]
    fi.remember_mapping(db, headers, {"question": 2, "option_a": 0})
    db.commit()
    mapping, row = fi.recall_mapping(db, headers)
    assert mapping == {"question": 2, "option_a": 0} and row.uses == 1
    reordered = [headers[3], headers[2], headers[1], headers[0]]                                    # same headers, other order: same signature
    assert fi.recall_mapping(db, reordered)[0] == {"question": 1, "option_a": 3}
    fi.remember_mapping(db, headers, {"question": 1})
    db.commit()
    assert fi.recall_mapping(db, headers)[1].uses == 2
    assert fi.recall_mapping(db, headers + ["another"]) == ({}, None)


def test_the_posted_mapping_is_read_safely():
    assert fi.read_mapping_form({"map_question": "1", "map_answer": "", "map_topic": "9", "map_number": "x", "other": "1"}, 3) == {"question": 1}


@pytest.mark.parametrize("raw,expected", [("b", "B"), ("B", "B"), ("(b)", "B"), ("B.", "B"), ("Option B", "B"), ("answer: c", "C"), ("[d]", "D"),
                                          ("1", None), ("4", None), ("AB", None), ("E", None), ("none", None), ("", None)])
def test_answers_are_letters_a_to_d_only(raw, expected):
    assert fi.normalise_answer(raw) == expected


# --------------------------------------------------------------------------- turning rows into questions

def parse(rows, headers=HEADERS, mapping=None, **kw):
    table = fi.read_table(csv_bytes(rows, headers), "q.csv")
    return fi.parse_rows(table, mapping if mapping is not None else fi.guess_mapping(table.headers), SUBJECTS, kw.get("topics"))


def test_a_clean_sheet_becomes_questions():
    parsed = parse(rows_of(3))
    assert parsed.ok and [q["number"] for q in parsed.questions] == [1, 2, 3] and parsed.errors == [] and parsed.warnings == []
    q = parsed.questions[0]
    assert q["answer"] == "B" and q["options"] == ["alpha 1", "beta 1", "gamma 1", "delta 1"] and q["subject"] == "Polity" and q["explanation"] == "Because 1."
    from app import json_import
    assert q["hash"] == json_import.norm_hash(q["text"], q["options"])


def test_required_columns_and_distinct_columns_are_checked():
    table = fi.read_table(csv_bytes(rows_of(1)), "q.csv")
    parsed = fi.parse_rows(table, {"question": 1, "option_a": 2}, SUBJECTS)
    assert not parsed.ok and "Option B" in parsed.errors[0] and "Option D" in parsed.errors[0]
    clash = fi.parse_rows(table, {"question": 1, "option_a": 2, "option_b": 2, "option_c": 4, "option_d": 5}, SUBJECTS)
    assert "same column" in clash.errors[0]


def test_every_bad_row_is_named_and_blocks_the_import():
    rows = rows_of(6)
    rows[0][0] = "x"                                     # no number
    rows[1][1] = ""                                      # no question
    rows[2][3] = ""                                      # empty option B
    rows[4][0] = "4"                                     # 4 used twice (rows 4 and 5)
    parsed = parse(rows)
    text = " ".join(parsed.errors)
    assert not parsed.ok
    assert "line 2" in text and "line 3" in text and "line 4: option B is empty" in text and "used twice: 4 (lines" in text


def test_answers_that_are_not_letters_are_flagged_not_guessed():
    rows = rows_of(4)
    rows[0][6], rows[1][6], rows[2][6] = "3", "", "c"
    parsed = parse(rows)
    assert parsed.ok
    answers = {q["number"]: (q["answer"], q["flags"]) for q in parsed.questions}
    assert answers[1] == (None, ["answer_unclear"]) and answers[2] == (None, ["no_answer_found"]) and answers[3] == ("C", [])
    assert any("can't be read as A, B, C or D" in w and "digits like 1–4 are never guessed" in w for w in parsed.warnings)
    assert any("No answer for question 2" in w for w in parsed.warnings)


def test_without_an_answer_column_every_question_is_flagged():
    parsed = parse(rows_of(2), mapping={"number": 0, "question": 1, "option_a": 2, "option_b": 3, "option_c": 4, "option_d": 5})
    assert parsed.ok and all(q["answer"] is None and "no_answer_found" in q["flags"] for q in parsed.questions)
    assert any("No answer column" in w for w in parsed.warnings)


def test_subjects_outside_the_fixed_list_are_left_blank_and_topics_need_to_exist():
    rows = rows_of(2)
    rows[0][8], rows[1][8] = "Astrology", "history"
    table = fi.read_table(csv_bytes([r + ["Ordinance", ] for r in rows], HEADERS + ["Topic"]), "q.csv")
    parsed = fi.parse_rows(table, fi.guess_mapping(table.headers), SUBJECTS, {"History": {"ordinance"}})
    assert [q["subject"] for q in parsed.questions] == [None, "History"] and [q["topic"] for q in parsed.questions] == [None, "Ordinance"]
    assert any("Astrology" in w for w in parsed.warnings)


def test_without_a_number_column_questions_are_numbered_in_order_and_gaps_are_reported_otherwise():
    parsed = parse(rows_of(3), mapping={"question": 1, "option_a": 2, "option_b": 3, "option_c": 4, "option_d": 5, "answer": 6})
    assert [q["number"] for q in parsed.questions] == [1, 2, 3]
    rows = rows_of(3)
    rows[2][0] = "6"
    gaps = parse(rows)
    assert any("missing from the file: 3, 4, 5" in w for w in gaps.warnings)
    assert parse([["Q 7.", stem(), "a", "b", "c", "d", "A", "", ""]]).questions[0]["number"] == 7      # "Q 7." is read as 7


def test_hindi_and_oddities_get_the_usual_flags():
    rows = rows_of(1)
    rows[0][1] = "भारत की राजधानी क्या है?"
    assert "hindi_text" in parse(rows).questions[0]["flags"]


# --------------------------------------------------------------------------- Word documents

def test_a_document_with_answers_under_each_question_is_read():
    parsed = fi.parse_docx(docx_bytes(["Some instructions here."] + docx_questions(4)))
    assert parsed.ok and [q["number"] for q in parsed.questions] == [1, 2, 3, 4] and parsed.method == "docx" and parsed.answer_source == "inline"
    assert [q["answer"] for q in parsed.questions] == ["A", "B", "C", "D"]
    assert parsed.questions[0]["options"] == ["alpha", "beta", "gamma", "delta"] and not parsed.questions[0]["flags"]
    assert parsed.warnings == []


def test_a_document_without_answers_flags_every_question():
    parsed = fi.parse_docx(docx_bytes(docx_questions(3, answers=False)))
    assert parsed.ok and all(q["answer"] is None and "no_answer_found" in q["flags"] for q in parsed.questions)
    assert any("has no answers" in w for w in parsed.warnings)


def test_a_key_at_the_end_of_a_document_fills_the_answers():
    lines = docx_questions(6, answers=False) + ["ANSWER KEY", "1-b 2-d 3-a 4-c 5-b 6-a"]
    parsed = fi.parse_docx(docx_bytes(lines))
    assert [q["answer"] for q in parsed.questions] == ["B", "D", "A", "C", "B", "A"] and parsed.answer_source == "paper_end"
    assert any("gave 6 answers" in w for w in parsed.warnings)
    assert all("no_answer_found" not in q["flags"] for q in parsed.questions)


def test_tables_in_a_document_are_read_row_by_row():
    lines = [f"1. {stem()}", "Match the following."]
    table = [["List I", "List II"], ["A. Ordinance", "1. Article 123"]]
    doc = docx_bytes(lines + ["(a) 1 (b) 2", "(c) 3 (d) 4"], tables=[table])
    parsed = fi.parse_docx(doc)
    assert parsed.ok and len(parsed.questions) == 1


def test_pictures_in_a_document_are_reported_not_imported():
    parsed = fi.parse_docx(docx_bytes(docx_questions(2), picture=True))
    assert parsed.ok and len(parsed.questions) == 2
    assert any("1 picture" in w and "aren't imported" in w for w in parsed.warnings)


def test_missing_numbers_in_a_document_are_listed():
    lines = docx_questions(3)
    parsed = fi.parse_docx(docx_bytes(lines[:4] + [f"3. {stem()}", "(a) a (b) b", "(c) c (d) d", "Answer: (a)"]))         # 1, then 3: 2 is missing
    assert parsed.ok and any("look like two questions merged" in w or "not detected" in w for w in parsed.warnings)

def test_a_document_with_no_numbered_questions_is_refused_and_a_non_docx_is_an_error():
    parsed = fi.parse_docx(docx_bytes(["Just a paragraph of prose.", "Another one."], numbered=3))
    assert not parsed.ok and "numbered" in parsed.errors[0] and any("automatic numbering" in w for w in parsed.warnings)
    with pytest.raises(fi.FileImportError):
        fi.parse_docx(b"this is not a docx")


# --------------------------------------------------------------------------- pictures

def test_pictures_are_joined_into_one_pdf_page_each(tmp_path):
    import pdfplumber
    out = tmp_path / "pages.pdf"
    assert fi.images_to_pdf([png("white"), png("black", (100, 100)), png("gray")], str(out)) == 3
    with pdfplumber.open(out) as pdf:
        assert len(pdf.pages) == 3
    with pytest.raises(fi.FileImportError):
        fi.images_to_pdf([png(), b"not an image"], str(tmp_path / "bad.pdf"))
    with pytest.raises(fi.FileImportError):
        fi.images_to_pdf([], str(tmp_path / "none.pdf"))


# --------------------------------------------------------------------------- saving

def new_paper(db):
    paper = models.Paper(title=f"File paper {next(_run)}", exam_type=models.ExamType.full_length, status="ready")
    db.add(paper)
    db.commit()
    return paper


def test_saving_creates_needs_review_questions_with_their_source(db):
    paper = new_paper(db)
    parsed = parse(rows_of(3))
    parsed.method = "xlsx"
    counts = fi.save_questions(db, None, parsed, paper)
    db.commit()
    assert counts["created"] == 3 and counts["skipped_existing"] == 0 and "duplicates" in counts
    q = db.query(models.Question).filter_by(paper_id=paper.id, question_number=1).one()
    assert (q.status, q.source, q.extraction_method, q.answer_source, q.explanation_status) == (QStatus.NEEDS_REVIEW, "xlsx", "xlsx", "file", "unverified")
    assert q.correct_answer == "B" and q.subject_id is not None and q.norm_hash
    again = fi.save_questions(db, None, parse(rows_of(2)), paper)                                          # numbers 1-2 exist
    assert again["created"] == 0 and again["skipped_existing"] == 2


# --------------------------------------------------------------------------- the pages

def post_file(admin, name, data, mime="application/octet-stream", target="new", extra_files=()):
    files = [("files", (name, data, mime))] + [("files", f) for f in extra_files]
    return admin.post("/admin/import/file/read", data={"target": target}, files=files)


def token_of(response):
    m = re.search(r'name="token" value="([0-9a-f]{32})"', response.text)
    assert m, response.text[:600]
    return m.group(1)


def apply(admin, token, title=None, mapping=None, **extra):
    data = {"token": token, "title": title or f"File import {next(_run)}", "exam_type": "full_length", "source_type": "coaching_test",
            "marks_per_question": "2", "negative_fraction": "1/3", **extra}
    for field, index in (mapping or {}).items():
        data["map_" + field] = str(index)
    return admin.post("/admin/import/file/apply", data=data)


def paper_named(db, title):
    db.rollback()
    return db.query(models.Paper).filter_by(title=title).one()


def test_the_upload_page_and_menu_link(admin):
    page = admin.get("/admin/import/file").text
    assert "Import questions from a file" in page and ".docx" in page
    assert 'href="/admin/import/file"' in admin.get("/admin").text


def test_a_csv_is_previewed_with_a_guessed_mapping_then_imported(admin, db):
    r = post_file(admin, "sheet.csv", csv_bytes(rows_of(3)), "text/csv")
    assert r.status_code == 200
    page = r.text
    assert "Which column is which?" in page and "Column 2: Question" in page and "No problems" in page and "3 questions can be imported" in page
    assert re.search(r'<option value="1" selected>Column 2: Question', page)
    token = token_of(r)
    title = f"CSV paper {next(_run)}"
    r = apply(admin, token, title, mapping=fi.guess_mapping(HEADERS))
    assert r.status_code == 303 and re.fullmatch(r"/review/\d+", r.headers["location"])
    paper = paper_named(db, title)
    qs = db.query(models.Question).filter_by(paper_id=paper.id).all()
    assert len(qs) == 3 and all(q.status == QStatus.NEEDS_REVIEW and q.source == "csv" and q.answer_source == "file" for q in qs)
    assert paper.status == "ready" and paper.key_source == "Imported table file — unverified" and paper.marks_per_question == 2
    log = db.query(models.AuditLog).filter_by(paper_id=paper.id, action="paper.file_import").one()
    assert '"kind": "table"' in log.detail_json and '"created": 3' in log.detail_json
    review = admin.get(r.headers["location"]).text
    assert "Imported 3 questions from sheet.csv" in review and "no original page to compare with" in review


def test_the_mapping_is_remembered_for_the_next_file_with_the_same_headers(admin, db):
    headers = ["Nr", "Stem", "P", "Q", "R", "S", "Right"]                                                # nothing here can be guessed
    rows = [[str(i), stem(), "a", "b", "c", "d", "A"] for i in range(1, 3)]
    first = post_file(admin, "one.csv", csv_bytes(rows, headers))
    assert "Using the choice you made last time" not in first.text
    mapping = {"number": 0, "question": 1, "option_a": 2, "option_b": 3, "option_c": 4, "option_d": 5, "answer": 6}
    assert apply(admin, token_of(first), mapping=mapping).status_code == 303
    rows2 = [[str(i), stem(), "a", "b", "c", "d", "C"] for i in range(1, 4)]
    second = post_file(admin, "two.csv", csv_bytes(rows2, headers))
    assert "Using the choice you made last time" in second.text and "3 questions can be imported" in second.text
    assert re.search(r'<option value="6" selected>Column 7: Right', second.text)
    db.rollback()
    assert db.query(models.ImportMapping).filter_by(signature=fi.signature(headers)).one().uses == 1


def test_an_import_with_problems_is_refused_and_saves_nothing(admin, db):
    rows = rows_of(3)
    rows[1][1] = ""
    r = post_file(admin, "bad.csv", csv_bytes(rows))
    assert "the import can&#39;t go ahead" in r.text or "the import can't go ahead" in r.text
    assert "disabled" in r.text.split("Import 2 question")[0].rsplit("<button", 1)[1]
    title = f"Bad CSV {next(_run)}"
    r = apply(admin, token_of(r), title, mapping=fi.guess_mapping(HEADERS))
    assert r.status_code == 400 and "problems listed below" in r.text
    db.rollback()
    assert db.query(models.Paper).filter_by(title=title).count() == 0


def test_changing_the_mapping_changes_the_result(admin, db):
    r = post_file(admin, "map.csv", csv_bytes(rows_of(2)))
    token = token_of(r)
    wrong = {"question": 0, "option_a": 2, "option_b": 3, "option_c": 4, "option_d": 5}              # the number column as the question
    preview = admin.post("/admin/import/file/preview", data={"token": token, **{"map_" + k: str(v) for k, v in wrong.items()}})
    assert preview.status_code == 200 and "2 questions can be imported" in preview.text
    missing = admin.post("/admin/import/file/preview", data={"token": token, "map_question": "1"})
    assert "Choose a column for" in missing.text and "Option A" in missing.text


def test_an_xlsx_with_two_sheets_can_switch_sheet(admin, db):
    good = [HEADERS] + rows_of(2)
    other = [HEADERS] + rows_of(4)
    r = post_file(admin, "book.xlsx", xlsx_bytes({"Prelims": good, "Mains": other}))
    assert "2 questions can be imported" in r.text and "Use this sheet" in r.text
    token = token_of(r)
    switched = admin.post("/admin/import/file/preview", data={"token": token, "sheet": "Mains", "switch_sheet": "1"})
    assert "4 questions can be imported" in switched.text
    title = f"XLSX paper {next(_run)}"
    assert apply(admin, token, title, mapping=fi.guess_mapping(HEADERS), sheet="Mains").status_code == 303
    paper = paper_named(db, title)
    assert db.query(models.Question).filter_by(paper_id=paper.id, source="xlsx").count() == 4


def test_a_document_is_previewed_and_imported(admin, db):
    r = post_file(admin, "test.docx", docx_bytes(docx_questions(3)))
    assert r.status_code == 200 and "Document report" in r.text and "3 questions can be imported" in r.text and "Which column is which" not in r.text
    title = f"DOCX paper {next(_run)}"
    assert apply(admin, token_of(r), title).status_code == 303
    paper = paper_named(db, title)
    qs = {q.question_number: q for q in db.query(models.Question).filter_by(paper_id=paper.id)}
    assert sorted(qs) == [1, 2, 3] and qs[1].source == "docx" and qs[1].answer_source == "inline" and qs[2].correct_answer == "B"


def test_adding_to_an_existing_paper_only_fills_new_numbers(admin, db):
    paper = new_paper(db)
    db.add(models.Question(paper_id=paper.id, question_number=1, text="Already here?", option_a="a", option_b="b", option_c="c", option_d="d",
                           status=QStatus.VERIFIED, source="pdf_text"))
    db.commit()
    r = post_file(admin, "more.csv", csv_bytes(rows_of(3)), target=str(paper.id))
    assert f"Adding to “{paper.title}”" in r.text
    assert apply(admin, token_of(r), mapping=fi.guess_mapping(HEADERS)).status_code == 303
    db.rollback()
    qs = {q.question_number: q for q in db.query(models.Question).filter_by(paper_id=paper.id)}
    assert sorted(qs) == [1, 2, 3] and qs[1].text == "Already here?" and qs[1].status == QStatus.VERIFIED and qs[2].status == QStatus.NEEDS_REVIEW


def test_the_same_file_twice_is_refused_unless_you_say_so(admin, db):
    data = csv_bytes(rows_of(2))
    first = post_file(admin, "same.csv", data)
    assert apply(admin, token_of(first), mapping=fi.guess_mapping(HEADERS)).status_code == 303
    second = post_file(admin, "same.csv", data)
    refused = apply(admin, token_of(second), mapping=fi.guess_mapping(HEADERS))
    assert refused.status_code == 400 and "already imported" in refused.text
    third = post_file(admin, "same.csv", data)
    assert apply(admin, token_of(third), mapping=fi.guess_mapping(HEADERS), allow_duplicate="true").status_code == 303
    db.rollback()
    found = db.query(models.QuestionDuplicate).filter_by(status="open").count()
    assert found >= 2                                                                                     # the second import repeats every question


def test_the_upload_step_refuses_what_it_cannot_read(admin, db):
    assert "Choose a file" in admin.post("/admin/import/file/read", data={"target": "new"}, files=[("files", ("", b"", "text/plain"))]).text
    assert post_file(admin, "paper.pdf", b"%PDF").status_code == 400
    assert "isn't a kind of file" in post_file(admin, "paper.pdf", b"%PDF").text or "isn&#39;t a kind of file" in post_file(admin, "paper.pdf", b"%PDF").text
    assert post_file(admin, "a.csv", csv_bytes(rows_of(1)), extra_files=[("b.docx", docx_bytes(docx_questions(1)), "application/octet-stream")]).status_code == 400
    assert post_file(admin, "a.csv", csv_bytes(rows_of(1)), extra_files=[("b.csv", csv_bytes(rows_of(1)), "text/csv")]).status_code == 400
    assert post_file(admin, "empty.csv", b"").status_code == 400
    assert post_file(admin, "broken.xlsx", b"nope").status_code == 400
    assert post_file(admin, "p.png", png(), target=str(new_paper(db).id)).status_code == 400
    assert post_file(admin, "none.docx", b"not a docx").status_code == 400


def test_bad_or_expired_tokens_are_404(admin):
    for token in ("../../etc", "0" * 31, "z" * 32, "", "0" * 32):
        assert admin.post("/admin/import/file/preview", data={"token": token}).status_code == 404
        assert admin.post("/admin/import/file/apply", data={"token": token, "title": "x"}).status_code == 404


def test_pictures_become_a_scanned_paper_read_in_the_background(admin, db, monkeypatch):
    called = []
    monkeypatch.setattr(ocr_extract, "tesseract_cmd", lambda: "tesseract")
    monkeypatch.setattr(ingest, "process_paper", lambda *a, **k: called.append(a))
    r = post_file(admin, "page1.png", png("white"), "image/png", extra_files=[("page2.png", png("gray"), "image/png")])
    assert r.status_code == 200 and "Pictures report" in r.text and "2 pictures" in r.text and "Read 2 pictures as a paper" in r.text
    title = f"Picture paper {next(_run)}"
    resp = apply(admin, token_of(r), title)
    assert resp.status_code == 303
    paper = paper_named(db, title)
    assert paper.status == "processing" and paper.source_pdf_path.endswith(".pdf")
    import os
    import pdfplumber
    with pdfplumber.open(paper.source_pdf_path) as pdf:
        assert len(pdf.pages) == 2
    assert os.path.exists(paper.source_pdf_path) and called == [(paper.id, None, "auto")]
    assert '"pages": 2' in db.query(models.AuditLog).filter_by(paper_id=paper.id, action="paper.file_import").one().detail_json


def test_pictures_need_tesseract_and_nothing_is_created_without_it(admin, db, monkeypatch):
    def missing():
        raise ocr_extract.OcrUnavailable("Tesseract isn't installed.")
    monkeypatch.setattr(ocr_extract, "tesseract_cmd", missing)
    r = post_file(admin, "page.png", png(), "image/png")
    title = f"No OCR {next(_run)}"
    resp = apply(admin, token_of(r), title)
    assert resp.status_code == 400 and "Tesseract" in resp.text
    db.rollback()
    assert db.query(models.Paper).filter_by(title=title).count() == 0


def test_the_labels_for_imported_answers_and_explanations(db):
    q = models.Question(source="csv", explanation="x", explanation_status="unverified")
    assert attempts.explanation_label(q) == "Unverified (imported from a file)"
    assert attempts.explanation_label(models.Question(source="ai_json", explanation_status="unverified")) == "AI-supplied, unverified"
    from app.routes.keys import SOURCE_LABELS
    assert SOURCE_LABELS["file"] == "an imported file"


def test_the_file_import_pages_are_admin_only(admin, db, anon):
    token = token_of(post_file(admin, "priv.csv", csv_bytes(rows_of(1))))
    student = make_student_client(db, "filestudent")
    calls = [("get", "/admin/import/file", {}), ("post", "/admin/import/file/read", {"data": {"target": "new"}}),
             ("post", "/admin/import/file/preview", {"data": {"token": token}}), ("post", "/admin/import/file/apply", {"data": {"token": token, "title": "x"}})]
    for client, expected in ((student, 403), (anon, 303)):
        for method, url, kw in calls:
            assert getattr(client, method)(url, **kw).status_code == expected, url
    db.rollback()
    assert db.query(models.Paper).filter_by(title="x").count() == 0
